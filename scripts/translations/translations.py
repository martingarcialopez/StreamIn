#!/usr/bin/env python3
"""Translation review: String Catalogs ⇄ one spreadsheet, a sheet per screen.

Usage (through scripts/translations.sh, which provides openpyxl):
  translations.sh export [--out <workbook.numbers|.xlsx>] [--force]
  translations.sh import [<workbook.numbers|.xlsx>] [--dry-run]
  translations.sh preview <StreamIn.app> [<workbook.numbers|.xlsx>]

The workbook defaults to scripts/streamin-translations.numbers (.xlsx without
Numbers). preview is what translator-preview.sh runs to show the edits in the
simulator before they're imported.

export writes every on-screen string, grouped by the screen that shows it
(screens.json maps the Swift file that uses a key to a screen), with a column
per language that the reviewer edits in place and hidden copies of the
exported text. On a Mac with Numbers it writes a .numbers file by default, so
the reviewer edits it and saves it in place; Numbers itself converts it
(numbers.applescript). An existing workbook is never replaced without --force.

import writes every edited cell back to its catalog, keeping each catalog's
JSON style, and regenerates the package symbols. Nothing is written while any
edit has an error (placeholders that differ from the current text, an empty
cell, a key that no longer exists); edits whose text changed in the code since
the export are skipped and listed, so they can be reviewed again.
"""

import argparse
import json
import plistlib
import re
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SCREENS = Path(__file__).resolve().with_name("screens.json")
NUMBERS_SCRIPT = Path(__file__).resolve().with_name("numbers.applescript")
NUMBERS_APP = Path("/Applications/Numbers.app")
WORKBOOK_NAME = "streamin-translations"

LANGUAGE_NAMES = {"en": "Inglés", "es": "Español"}
KEY_HEADER = "Clave"
CONTEXT_HEADER = "Contexto"
PLACEHOLDERS_HEADER = "Variables"
NOTES_HEADER = "Notas"
CATALOG_HEADER = "Catálogo"
ORIGINAL_SUFFIX = " (original)"
GUIDE_SHEET = "Cómo usar esta hoja"

# printf-style specifiers String Catalogs use: %@, %lld, %d, %f, positional %2$@.
PLACEHOLDER = re.compile(r"%(?:(\d+)\$)?(@|lld|ld|d|f|%)")
PLACEHOLDER_KINDS = {"@": "texto", "lld": "número", "ld": "número", "d": "número", "f": "decimal"}
IDENTIFIER = re.compile(r"[a-z][A-Za-z0-9]*")


# MARK: - Catalogs


@dataclass
class Catalog:
    path: Path
    raw: str
    data: dict
    sources: list  # Swift files of the module the catalog belongs to

    @property
    def strings(self) -> dict:
        return self.data["strings"]

    @property
    def is_package(self) -> bool:
        return self.path.parent.name == "Translations"

    def value(self, key: str, language: str):
        return self.strings[key].get("localizations", {}).get(language, {}).get("stringUnit", {}).get("value")

    def has_variants(self, key: str) -> bool:
        """Plural or device variations, or substitutions: more than one text per language."""
        entry = self.strings[key]
        return any("variations" in loc or "substitutions" in loc for loc in entry.get("localizations", {}).values())

    def set_value(self, key: str, language: str, text: str) -> None:
        localizations = self.strings[key].setdefault("localizations", {})
        localizations.setdefault(language, {})["stringUnit"] = {"state": "translated", "value": text}
        self.strings[key]["localizations"] = dict(sorted(localizations.items()))

    def serialized(self) -> str:
        separators = (",", " : ") if '" : ' in self.raw else (",", ": ")
        output = json.dumps(self.data, indent=2, ensure_ascii=False, separators=separators)
        return output + "\n" if self.raw.endswith("\n") else output


def load_catalogs(repo: Path) -> list:
    catalogs = []
    for path in sorted(repo.glob("Packages/*/Sources/*/Translations/*.xcstrings")):
        module = path.parent.parent
        catalogs.append(_catalog(path, module))
    for path in sorted(repo.glob("StreamIn/**/*.xcstrings")):
        catalogs.append(_catalog(path, repo / "StreamIn"))
    return catalogs


def _catalog(path: Path, module: Path) -> Catalog:
    raw = path.read_text(encoding="utf-8")
    sources = sorted(p for p in module.rglob("*.swift") if not p.name.startswith("GeneratedStringSymbols"))
    return Catalog(path, raw, json.loads(raw), sources)


def languages_of(catalogs: list) -> list:
    found = {lang for c in catalogs for e in c.strings.values() for lang in e.get("localizations", {})}
    source = catalogs[0].data.get("sourceLanguage", "en") if catalogs else "en"
    return [source] + sorted(found - {source})


def placeholders(text: str) -> list:
    """Specifiers as (position, kind), so "%@ %lld" equals "%1$@ %2$lld"."""
    found, implicit = [], 0
    for match in PLACEHOLDER.finditer(text or ""):
        position, kind = match.groups()
        if kind == "%":
            continue
        implicit += 1
        kind = "lld" if kind == "ld" else kind
        found.append((int(position) if position else implicit, kind))
    return sorted(found)


def describe_placeholders(text: str) -> str:
    names = []
    for match in PLACEHOLDER.finditer(text or ""):
        if match.group(2) != "%":
            names.append(f"{match.group(0)} = {PLACEHOLDER_KINDS[match.group(2)]}")
    return "\n".join(dict.fromkeys(names))


# MARK: - Screens


@dataclass
class Row:
    screen: int
    order: tuple
    catalog: Catalog
    key: str
    also_on: list = field(default_factory=list)


def rows_by_screen(repo: Path, catalogs: list, screens: dict):
    """Each used key, placed on the first screen whose files use it.

    Returns (rows, keys left out with the reason). Exits if a key is used in a
    Swift file that screens.json doesn't map, so a new view can't drop out of
    the review. Keys with plural forms or substitutions are left out: one cell
    can't hold them, so they're edited in Xcode.
    """
    file_screen = {}
    for index, screen in enumerate(screens["screens"]):
        for position, name in enumerate(screen.get("files", [])):
            file_screen[name] = (index, position)
    key_screen = {key: index for index, s in enumerate(screens["screens"]) for key in s.get("keys", [])}
    skip = screens.get("skipKeys", {})

    rows, unused, unmapped = [], [], {}
    for catalog in catalogs:
        texts = {path: path.read_text(encoding="utf-8") for path in catalog.sources}
        for key in catalog.strings:
            if key in skip:
                continue
            if catalog.has_variants(key):
                unused.append(f"{catalog.path.relative_to(repo).as_posix()}: {key} (plural forms; edit it in Xcode)")
                continue
            if key in key_screen:
                rows.append(Row(key_screen[key], (999, 0), catalog, key))
                continue
            pattern = re.compile(r"\." + re.escape(key) + r"\b" if IDENTIFIER.fullmatch(key) else '"' + re.escape(key) + '"')
            places = []
            for path, text in texts.items():
                match = pattern.search(text)
                if not match:
                    continue
                if path.name not in file_screen:
                    unmapped.setdefault(path.relative_to(repo).as_posix(), []).append(key)
                    continue
                index, position = file_screen[path.name]
                places.append((index, position, match.start()))
            if not places:
                if not any(key in keys for keys in unmapped.values()):
                    unused.append(f"{catalog.path.relative_to(repo).as_posix()}: {key} (no Swift file uses it)")
                continue
            places.sort()
            index, position, offset = places[0]
            others = sorted({screens["screens"][i]["sheet"] for i, _, _ in places} - {screens["screens"][index]["sheet"]})
            rows.append(Row(index, (position, offset), catalog, key, others))

    if unmapped:
        lines = "\n".join(f"  {path}: {', '.join(keys)}" for path, keys in sorted(unmapped.items()))
        sys.exit(f"These Swift files show strings but no screen in screens.json lists them:\n{lines}\n"
                 "Add each file to the screen that shows it.")
    rows.sort(key=lambda r: (r.screen, r.order))
    return rows, unused


# MARK: - Export


def export(repo: Path, out: Path, screens_path: Path = SCREENS):
    """Writes the workbook; returns (rows, keys left out and why)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    catalogs = load_catalogs(repo)
    screens = json.loads(screens_path.read_text(encoding="utf-8"))
    languages = languages_of(catalogs)
    rows, unused = rows_by_screen(repo, catalogs, screens)

    # The reviewer edits the language columns in place; the hidden originals tell
    # the import what changed, and whether the code changed since the export.
    names = [LANGUAGE_NAMES.get(language, language) for language in languages]
    hidden = [name + ORIGINAL_SUFFIX for name in names] + [CATALOG_HEADER]
    headers = [KEY_HEADER, CONTEXT_HEADER, PLACEHOLDERS_HEADER] + names + [NOTES_HEADER] + hidden
    widths = {KEY_HEADER: 26, CONTEXT_HEADER: 48, PLACEHOLDERS_HEADER: 16, NOTES_HEADER: 32}

    bold = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="3A3A3C")
    editable = PatternFill("solid", fgColor="FFF4C2")
    muted = Font(color="6E6E73")
    wrap = Alignment(wrap_text=True, vertical="top")

    workbook = Workbook()
    _write_guide(workbook.active, screens, languages)

    for index, screen in enumerate(screens["screens"]):
        screen_rows = [r for r in rows if r.screen == index]
        if not screen_rows:
            continue
        sheet = workbook.create_sheet(screen["sheet"])
        sheet.append(headers)
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column)
            cell.font, cell.fill, cell.alignment = bold, header_fill, wrap
            letter = cell.column_letter
            sheet.column_dimensions[letter].width = widths.get(header, 40)
        for row in screen_rows:
            source_text = row.catalog.value(row.key, languages[0]) or ""
            context = row.catalog.strings[row.key].get("comment", "")
            if row.also_on:
                context += f"\n(También en: {', '.join(row.also_on)})"
            texts = [row.catalog.value(row.key, language) or "" for language in languages]
            catalog_name = row.catalog.path.relative_to(repo).as_posix()
            sheet.append([row.key, context.strip(), describe_placeholders(source_text)]
                         + texts + [None] + texts + [catalog_name])
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                header = headers[cell.column - 1]
                cell.alignment = wrap
                if header in names or header == NOTES_HEADER:
                    cell.fill = editable
                elif header in (KEY_HEADER, PLACEHOLDERS_HEADER):
                    cell.font = muted
        for header in hidden:
            letter = sheet.cell(row=1, column=headers.index(header) + 1).column_letter
            sheet.column_dimensions[letter].hidden = True
        sheet.freeze_panes = "B2"

    out.parent.mkdir(parents=True, exist_ok=True)
    if out.suffix == ".numbers":
        with numbers_temporary_directory(out.parent) as temp:
            xlsx = Path(temp) / f"{out.stem}.xlsx"
            workbook.save(xlsx)
            convert_with_numbers("to-numbers", xlsx, out)
    else:
        workbook.save(out)
    return rows, unused


def numbers_temporary_directory(beside: Path):
    """A visible folder next to the workbook, a place sandboxed Numbers reaches,
    unlike the system temp folder."""
    return tempfile.TemporaryDirectory(prefix="translations-tmp-", dir=beside)


def convert_with_numbers(mode: str, source: Path, destination: Path) -> None:
    """mode is "to-numbers" or "to-xlsx"; Numbers asks once for permission to be scripted."""
    result = subprocess.run(
        ["osascript", str(NUMBERS_SCRIPT), mode, str(source), str(destination), GUIDE_SHEET],
        capture_output=True, text=True,
    )
    if result.returncode != 0:
        sys.exit(f"Numbers couldn't convert {source.name}: {result.stderr.strip()}")


def _write_guide(sheet, screens: dict, languages: list) -> None:
    from openpyxl.styles import Alignment, Font

    names = " y ".join(LANGUAGE_NAMES.get(lang, lang).lower() for lang in languages)
    sheet.title = GUIDE_SHEET
    lines = [
        ("Revisión de textos de StreamIn", "title"),
        (f"Cada pestaña es una pantalla de la app, en el orden en que la usas. Cada fila es un texto, en {names}.", None),
        ("", None),
        ("Qué hacer", "heading"),
        ("• Si un texto se puede mejorar, cámbialo directamente en su celda (columnas amarillas).", None),
        ("• Si está bien, déjalo como está. No dejes ninguna celda de texto vacía.", None),
        ("• «Notas» es para dudas o comentarios; se leerán todas.", None),
        ("• No cambies las columnas blancas, el nombre de las pestañas ni el orden de las filas, y no añadas filas.", None),
        ("", None),
        ("Variables", "heading"),
        ("%@ y %lld son huecos que la app rellena: %@ con un texto (un título, un país), %lld con un número.", None),
        ("Tienen que seguir en el texto, tal cual. Si hay varios (%1$@, %2$@) puedes cambiarlos de orden.", None),
        ("", None),
        ("Contexto", "heading"),
        ("«Contexto» dice dónde aparece el texto y qué significa. Botones, pestañas y títulos tienen poco sitio: mejor cortos.", None),
        ("La app trata al usuario de tú.", None),
        ("Para ver tus cambios en la app, aunque no hayas guardado, ejecuta en la terminal, desde la carpeta del proyecto:", None),
        ("    scripts/translator-preview.sh es     (o en, para inglés)", None),
        ("", None),
        ("Pantallas", "heading"),
    ]
    lines += [(f"• {s['sheet']}: {s['description']}", None) for s in screens["screens"]]
    lines += [("", None), ("Cuando termines, guarda el fichero (⌘S) y devuélvelo tal cual.", "heading")]
    for text, style in lines:
        sheet.append([text])
        cell = sheet.cell(row=sheet.max_row, column=1)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            cell.font = Font(bold=True, size=16)
        elif style == "heading":
            cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 110


# MARK: - Import


@dataclass
class Report:
    changes: list = field(default_factory=list)
    stale: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    symbol_failures: list = field(default_factory=list)


@dataclass
class Edit:
    catalog: str  # path relative to the repo
    key: str
    language: str
    text: str


def read_workbook(repo: Path, workbook_path: Path):
    """Checks every edited cell against the catalogs.

    Returns (report, catalogs by relative path with the valid edits applied in
    memory, the valid edits). Nothing is written.
    """
    from openpyxl import load_workbook

    if workbook_path.suffix == ".numbers":
        # Read through Numbers: if the document is open there, as on screen,
        # unsaved edits included.
        with numbers_temporary_directory(workbook_path.parent) as temp:
            xlsx = Path(temp) / f"{workbook_path.stem}.xlsx"
            convert_with_numbers("to-xlsx", workbook_path, xlsx)
            return read_workbook(repo, xlsx)

    catalogs = {c.path.relative_to(repo).as_posix(): c for c in load_catalogs(repo)}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    report = Report()
    applied = []

    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers = [_text(h) for h in next(rows, [])]
        if KEY_HEADER not in headers:
            continue
        column = {h: i for i, h in enumerate(headers)}
        missing = [h for h in (CATALOG_HEADER, NOTES_HEADER) if h not in column]
        if missing:
            report.errors.append(f"{sheet.title}: missing column(s) {', '.join(missing)}")
            continue
        codes = {name: code for code, name in LANGUAGE_NAMES.items()}
        language_columns = {}
        for header in headers:
            name = header.removesuffix(ORIGINAL_SUFFIX)
            if header.endswith(ORIGINAL_SUFFIX) and name in column:
                language_columns[codes.get(name, name)] = (column[header], column[name])

        for number, values in enumerate(rows, start=2):
            cells = list(values) + [None] * (len(headers) - len(values))
            key, catalog_name = _text(cells[column[KEY_HEADER]]), _text(cells[column[CATALOG_HEADER]])
            if not key:
                continue
            where = f"{sheet.title}, row {number} ({key})"
            note = _text(cells[column[NOTES_HEADER]])
            if note:
                report.notes.append(f"{where}: {note}")
            edits = {
                language: (_text(cells[original]), _text(cells[edited]).replace("\r\n", "\n"))
                for language, (original, edited) in language_columns.items()
            }
            edits = {language: pair for language, pair in edits.items() if pair[0] != pair[1]}
            catalog = catalogs.get(catalog_name)
            if catalog is None or key not in catalog.strings:
                if edits:
                    report.errors.append(f"{where}: the key no longer exists in {catalog_name}")
                continue
            for language, (exported, edited) in sorted(edits.items()):
                if not edited:
                    report.errors.append(f"{where} [{language}]: the cell is empty; write the text back")
                    continue
                if catalog.has_variants(key):
                    report.errors.append(f"{where} [{language}]: it has plural forms now; edit it in Xcode")
                    continue
                live = catalog.value(key, language) or ""
                if exported != live:
                    report.stale.append(f"{where} [{language}]: the text changed since the export, now “{live}”")
                    continue
                if edited == live:
                    continue
                reference = live or catalog.value(key, languages_of([catalog])[0]) or ""
                if placeholders(edited) != placeholders(reference):
                    report.errors.append(
                        f"{where} [{language}]: “{edited}” must keep the variables of “{reference}”")
                    continue
                if live and len(edited) > max(len(live) * 1.5, len(live) + 10):
                    report.warnings.append(f"{where} [{language}]: {len(live)} → {len(edited)} characters, check it fits")
                catalog.set_value(key, language, edited)
                applied.append(Edit(catalog_name, key, language, edited))
                report.changes.append(f"{where} [{language}]: “{live}” → “{edited}”")
    return report, catalogs, applied


def import_workbook(repo: Path, workbook_path: Path, dry_run: bool = False, regenerate_symbols: bool = True) -> Report:
    report, catalogs, applied = read_workbook(repo, workbook_path)
    if report.errors or dry_run:
        return report
    touched = {edit.catalog for edit in applied}
    # Every catalog first, so a symbol failure never leaves half the changes unwritten.
    for name in sorted(touched):
        catalogs[name].path.write_text(catalogs[name].serialized(), encoding="utf-8")
    for name in sorted(touched):
        catalog = catalogs[name]
        if not (regenerate_symbols and catalog.is_package):
            continue
        result = subprocess.run(
            ["xcrun", "xcstringstool", "generate-symbols", str(catalog.path),
             "--output-directory", str(catalog.path.parent), "--language", "swift"],
            capture_output=True, text=True,
        )
        if result.returncode != 0:
            report.symbol_failures.append(f"{name}: {result.stderr.strip() or 'xcstringstool failed'}")
    return report


# MARK: - Preview


def preview_workbook(repo: Path, workbook_path: Path, app: Path) -> Report:
    """Writes the workbook's valid edits into a built StreamIn.app's compiled strings.

    The catalogs aren't touched, so the reviewer can try copy in the simulator
    before anything is imported. Cells with errors are left out of the app.
    """
    report, _, applied = read_workbook(repo, workbook_path)
    for edit in applied:
        path = compiled_strings(app, edit.catalog, edit.language)
        if not path.exists():
            report.warnings.append(f"{edit.key} [{edit.language}]: no {path.relative_to(app)} in the app")
            continue
        with path.open("rb") as file:
            strings = plistlib.load(file)
        strings[edit.key] = edit.text
        with path.open("wb") as file:
            plistlib.dump(strings, file, fmt=plistlib.FMT_BINARY)
    return report


def compiled_strings(app: Path, catalog: str, language: str) -> Path:
    """Where Xcode compiles a catalog's strings for one language inside the .app."""
    parts = Path(catalog).parts
    lproj = f"{language}.lproj/{Path(catalog).stem}.strings"
    if parts[0] == "Packages":
        return app / f"{parts[1]}_{parts[3]}.bundle" / lproj
    return app / lproj


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def print_report(report: Report, dry_run: bool, preview: bool = False) -> None:
    if preview:
        errors_title, changes_title = "Not shown — fix these cells", "Shown in the app"
    else:
        errors_title = "Errors — fix these cells, nothing was written"
        changes_title = "Changes" + (" (dry run, nothing written)" if dry_run or report.errors else "")
    sections = [
        (errors_title, report.errors),
        (changes_title, report.changes),
        ("Skipped — the text changed since the export; review these again", report.stale),
        ("Check the length on screen", report.warnings),
        ("Reviewer's notes", report.notes),
        ("Catalogs written, but their symbols weren't regenerated — run the localize "
         "skill's xcstringstool command on them", report.symbol_failures),
    ]
    for title, items in sections:
        if items:
            print(f"\n{title} ({len(items)}):")
            for item in items:
                print(f"  • {item}")
    if not any(items for _, items in sections):
        print("No edits in the workbook.")


# MARK: - CLI


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)
    extension = ".numbers" if NUMBERS_APP.exists() else ".xlsx"
    default_workbook = REPO / "scripts" / f"{WORKBOOK_NAME}{extension}"
    workbook_help = f"a .numbers or .xlsx workbook (default: {default_workbook.relative_to(REPO)})"
    export_parser = commands.add_parser("export", help="Write the review workbook")
    export_parser.add_argument("--out", type=Path, default=default_workbook, help=workbook_help)
    export_parser.add_argument("--force", action="store_true", help="Replace an existing workbook")
    import_parser = commands.add_parser("import", help="Apply a reviewed workbook to the catalogs")
    import_parser.add_argument("workbook", type=Path, nargs="?", default=default_workbook, help=workbook_help)
    import_parser.add_argument("--dry-run", action="store_true", help="Report what would change, write nothing")
    preview_parser = commands.add_parser("preview", help="Put a workbook's edits into a built StreamIn.app")
    preview_parser.add_argument("app", type=Path, help="the .app to change (translator-preview.sh passes a copy)")
    preview_parser.add_argument("workbook", type=Path, nargs="?", default=default_workbook, help=workbook_help)
    args = parser.parse_args()

    if args.command == "preview":
        report = preview_workbook(REPO, args.workbook.resolve(), args.app.resolve())
        print_report(report, dry_run=False, preview=True)
    elif args.command == "export":
        out = args.out.resolve()
        if out.suffix not in (".numbers", ".xlsx"):
            sys.exit(f"--out must end in .numbers or .xlsx, got {out.name}")
        if out.exists() and not args.force:
            sys.exit(f"{out} already exists and may hold a review in progress.\n"
                     "Import it first, or pass --force to replace it (--out writes another file).")
        if out.exists():
            out.unlink()
        rows, unused = export(REPO, out)
        print(f"Wrote {len(rows)} strings on {len({r.screen for r in rows})} screens to {out}")
        if unused:
            print("Left out:\n  " + "\n  ".join(unused))
    else:
        report = import_workbook(REPO, args.workbook.resolve(), dry_run=args.dry_run)
        print_report(report, args.dry_run)
        if report.errors or report.symbol_failures:
            sys.exit(1)


if __name__ == "__main__":
    main()
