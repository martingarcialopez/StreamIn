"""Tests for translations.py: export → fill in proposals → import, on a tiny repo.

Run with scripts/translations.sh test (CI does, before the app build).
"""

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import translations as t

SCREENS = {
    "screens": [
        {"sheet": "Lista", "description": "La lista.", "files": ["ListView.swift"]},
        {"sheet": "General", "description": "Todo lo demás.", "files": ["RootView.swift"], "keys": ["CFBundleDisplayName"]},
    ],
    "skipKeys": {"CFBundleName": "never on screen"},
}


def entry(comment, en, es=None):
    localizations = {"en": {"stringUnit": {"state": "translated", "value": en}}}
    if es is not None:
        localizations["es"] = {"stringUnit": {"state": "translated", "value": es}}
    return {"comment": comment, "extractionState": "manual", "localizations": localizations}


class TranslationsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.repo = Path(self.temp.name)
        self.package_catalog = self.write_catalog(
            "Packages/List/Sources/List/Translations/Localizable.xcstrings",
            {
                "listTitle": entry("Title", "Watchlist", "Mi lista"),
                "listCount": entry("Count", "%lld movies by %@", "%lld películas de %@"),
                "listUnused": entry("Unused", "Old", "Viejo"),
                "listEmpty": entry("Empty", "Nothing here"),
            },
            separators=(",", ": "),
        )
        self.app_catalog = self.write_catalog(
            "StreamIn/Resources/Localizable.xcstrings",
            {"Settings": entry("Gear", "Settings", "Ajustes")},
            separators=(",", " : "),
        )
        self.write_catalog(
            "StreamIn/Resources/InfoPlist.xcstrings",
            {"CFBundleDisplayName": entry("Name", "StreamIn", "StreamIn"), "CFBundleName": entry("Name", "StreamIn")},
            separators=(",", " : "),
        )
        self.write("Packages/List/Sources/List/ListView.swift",
                   "Text(.listTitle)\nText(.listCount(3, name))\nText(.listEmpty)\n")
        self.write("StreamIn/App/RootView.swift", 'Button("Settings") {}\n')
        self.screens = self.repo / "screens.json"
        self.screens.write_text(json.dumps(SCREENS), encoding="utf-8")
        self.workbook = self.repo / "review.xlsx"

    def tearDown(self):
        self.temp.cleanup()

    def write(self, relative, text):
        path = self.repo / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        return path

    def write_catalog(self, relative, strings, separators):
        text = json.dumps({"sourceLanguage": "en", "strings": strings, "version": "1.0"},
                          indent=2, ensure_ascii=False, separators=separators) + "\n"
        return self.write(relative, text)

    def export(self):
        return t.export(self.repo, self.workbook, self.screens)

    def add_plural_key(self):
        catalog = json.loads(self.package_catalog.read_text(encoding="utf-8"))
        catalog["strings"]["listPlural"] = {"localizations": {"en": {"variations": {"plural": {
            "one": {"stringUnit": {"state": "translated", "value": "%lld movie"}},
            "other": {"stringUnit": {"state": "translated", "value": "%lld movies"}},
        }}}}}
        self.package_catalog.write_text(json.dumps(catalog, indent=2), encoding="utf-8")
        self.write("Packages/List/Sources/List/ListView.swift",
                   "Text(.listTitle)\nText(.listCount(3, name))\nText(.listEmpty)\nText(.listPlural(2))\n")

    def propose(self, key, **values):
        """Edits the language (es=…, en=…) or notes cells of the row for `key`, as the reviewer does."""
        book = load_workbook(self.workbook)
        for sheet in book.worksheets[1:]:
            headers = [c.value for c in sheet[1]]
            for row in sheet.iter_rows(min_row=2):
                if row[0].value != key:
                    continue
                for name, value in values.items():
                    header = t.NOTES_HEADER if name == "notes" else t.LANGUAGE_NAMES[name]
                    row[headers.index(header)].value = value
                book.save(self.workbook)
                return
        self.fail(f"{key} isn't in the workbook")

    def run_import(self, dry_run=False):
        return t.import_workbook(self.repo, self.workbook, dry_run=dry_run, regenerate_symbols=False)

    def value(self, path, key, language):
        return json.loads(path.read_text(encoding="utf-8"))["strings"][key]["localizations"][language]["stringUnit"]["value"]

    # MARK: Export

    def test_export_groups_keys_by_screen_in_source_order(self):
        self.export()
        book = load_workbook(self.workbook)
        self.assertEqual(book.sheetnames, [t.GUIDE_SHEET, "Lista", "General"])
        keys = {sheet.title: [r[0] for r in sheet.iter_rows(min_row=2, values_only=True)] for sheet in book.worksheets[1:]}
        self.assertEqual(keys["Lista"], ["listTitle", "listCount", "listEmpty"])
        self.assertEqual(keys["General"], ["Settings", "CFBundleDisplayName"])

    def test_export_explains_placeholders(self):
        self.export()
        sheet = load_workbook(self.workbook)["Lista"]
        row = next(r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] == "listCount")
        self.assertEqual(row[2], "%lld = número\n%@ = texto")

    def test_export_shows_one_editable_column_per_language_and_hides_the_originals(self):
        self.export()
        sheet = load_workbook(self.workbook)["Lista"]
        headers = [c.value for c in sheet[1]]
        visible = [h for h in headers if not sheet.column_dimensions[sheet.cell(row=1, column=headers.index(h) + 1).column_letter].hidden]
        self.assertEqual(visible, ["Clave", "Contexto", "Variables", "Inglés", "Español", "Notas"])
        row = next(r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] == "listTitle")
        self.assertEqual(row[headers.index("Español")], "Mi lista")
        self.assertEqual(row[headers.index("Español (original)")], "Mi lista")

    def test_export_leaves_out_keys_with_plural_forms(self):
        self.add_plural_key()
        rows, left_out = self.export()
        self.assertNotIn("listPlural", [r.key for r in rows])
        self.assertTrue(any("listPlural (plural forms" in line for line in left_out))

    def test_export_fails_when_a_file_with_strings_has_no_screen(self):
        self.write("Packages/List/Sources/List/NewView.swift", "Text(.listUnused)\n")
        with self.assertRaises(SystemExit) as raised:
            self.export()
        self.assertIn("NewView.swift: listUnused", str(raised.exception))

    # MARK: Import

    def test_import_writes_proposals_and_keeps_the_catalog_style(self):
        self.export()
        before = self.app_catalog.read_text(encoding="utf-8")
        self.propose("Settings", es="Configuración")
        self.propose("listTitle", en="My list")

        report = self.run_import()

        self.assertEqual(report.errors, [])
        self.assertEqual(len(report.changes), 2)
        self.assertEqual(self.value(self.app_catalog, "Settings", "es"), "Configuración")
        self.assertEqual(self.value(self.package_catalog, "listTitle", "en"), "My list")
        after = self.app_catalog.read_text(encoding="utf-8")
        changed = [(a, b) for a, b in zip(before.splitlines(), after.splitlines()) if a != b]
        self.assertEqual(changed, [('            "value" : "Ajustes"', '            "value" : "Configuración"')])

    def test_import_adds_a_missing_translation(self):
        self.export()
        self.propose("listEmpty", es="Aquí no hay nada")
        self.run_import()
        self.assertEqual(self.value(self.package_catalog, "listEmpty", "es"), "Aquí no hay nada")

    def test_import_accepts_reordered_positional_placeholders(self):
        self.export()
        self.propose("listCount", es="De %2$@: %1$lld películas")
        report = self.run_import()
        self.assertEqual(report.errors, [])
        self.assertEqual(self.value(self.package_catalog, "listCount", "es"), "De %2$@: %1$lld películas")

    def test_import_refuses_changed_placeholders_and_writes_nothing(self):
        self.export()
        self.propose("listTitle", es="Lista guardada")
        self.propose("listCount", es="%@ películas")
        before = self.package_catalog.read_text(encoding="utf-8")

        report = self.run_import()

        self.assertEqual(len(report.errors), 1)
        self.assertIn("listCount", report.errors[0])
        self.assertEqual(self.package_catalog.read_text(encoding="utf-8"), before)

    def test_import_skips_rows_changed_since_the_export(self):
        self.export()
        self.propose("listTitle", es="Guardadas")
        self.propose("Settings", es="Configuración")
        catalog = json.loads(self.package_catalog.read_text(encoding="utf-8"))
        catalog["strings"]["listTitle"]["localizations"]["es"]["stringUnit"]["value"] = "Tu lista"
        self.package_catalog.write_text(json.dumps(catalog, indent=2, ensure_ascii=False), encoding="utf-8")

        report = self.run_import()

        self.assertEqual(len(report.stale), 1)
        self.assertEqual(self.value(self.package_catalog, "listTitle", "es"), "Tu lista")
        self.assertEqual(self.value(self.app_catalog, "Settings", "es"), "Configuración")

    def test_import_reports_a_key_that_no_longer_exists(self):
        self.export()
        self.propose("listTitle", es="Guardadas")
        catalog = json.loads(self.package_catalog.read_text(encoding="utf-8"))
        del catalog["strings"]["listTitle"]
        self.package_catalog.write_text(json.dumps(catalog, indent=2), encoding="utf-8")

        report = self.run_import()

        self.assertEqual(len(report.errors), 1)
        self.assertIn("no longer exists", report.errors[0])

    def test_import_refuses_an_emptied_cell(self):
        self.export()
        self.propose("listTitle", es="")
        report = self.run_import()
        self.assertEqual(len(report.errors), 1)
        self.assertIn("empty", report.errors[0])
        self.assertEqual(self.value(self.package_catalog, "listTitle", "es"), "Mi lista")

    def test_import_refuses_a_key_that_gained_plural_forms(self):
        self.export()
        self.propose("listTitle", es="Guardadas")
        catalog = json.loads(self.package_catalog.read_text(encoding="utf-8"))
        catalog["strings"]["listTitle"]["localizations"]["es"] = {"variations": {"plural": {}}}
        self.package_catalog.write_text(json.dumps(catalog, indent=2, ensure_ascii=False), encoding="utf-8")

        report = self.run_import()

        self.assertEqual(len(report.errors), 1)
        self.assertIn("plural", report.errors[0])
        self.assertIn("variations", json.loads(self.package_catalog.read_text(encoding="utf-8"))
                      ["strings"]["listTitle"]["localizations"]["es"])

    def test_dry_run_writes_nothing(self):
        self.export()
        self.propose("Settings", es="Configuración")
        before = self.app_catalog.read_text(encoding="utf-8")
        report = self.run_import(dry_run=True)
        self.assertEqual(len(report.changes), 1)
        self.assertEqual(self.app_catalog.read_text(encoding="utf-8"), before)

    def test_import_collects_notes_and_warns_about_long_proposals(self):
        self.export()
        self.propose("Settings", es="Configuración de la aplicación", notes="¿Mejor «Ajustes»?")
        report = self.run_import()
        self.assertEqual(len(report.warnings), 1)
        self.assertIn("¿Mejor «Ajustes»?", report.notes[0])

    # MARK: Helpers

    def test_placeholders_treat_implicit_and_positional_alike(self):
        self.assertEqual(t.placeholders("%@ and %lld"), t.placeholders("%2$lld y %1$@"))
        self.assertEqual(t.placeholders("100%% %ld"), t.placeholders("%lld"))
        self.assertNotEqual(t.placeholders("%@"), t.placeholders("%lld"))
        self.assertNotEqual(t.placeholders("%d"), t.placeholders("%lld"))

    def test_real_catalogs_round_trip_byte_for_byte(self):
        for catalog in t.load_catalogs(t.REPO):
            self.assertEqual(catalog.serialized(), catalog.raw, catalog.path.name)


if __name__ == "__main__":
    unittest.main()
